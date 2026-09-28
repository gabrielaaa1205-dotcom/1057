import io
import datetime

from flask import Blueprint, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from db import q
from auth import login_required

bp = Blueprint("reports", __name__)


def to_xlsx(rows, filename, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    if not rows:
        ws.append(["Sin datos"])
    else:
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        for r in rows:
            ws.append([r[h] for h in headers])
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(40, len(h) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def respond(rows, filename):
    if request.args.get("format") == "xlsx":
        return to_xlsx(rows, filename)
    return jsonify(rows)


@bp.get("/reports/receptions-daily")
@login_required
def receptions_daily():
    date = request.args.get("date", datetime.date.today().isoformat())
    rows = q(
        """SELECT r.reception_number, c.name as cliente, r.reception_date as fecha, r.guide_number as guia,
                  r.container_number as contenedor, p.sku_code as sku, p.description as producto,
                  l.lot_code as lote, ri.qty_cases as cajas, ri.qty_units as unidades, ri.quality_status as estado_calidad
           FROM reception_items ri
           JOIN receptions r ON r.id = ri.reception_id
           JOIN clients c ON c.id = r.client_id
           JOIN products p ON p.id = ri.product_id
           LEFT JOIN lots l ON l.id = ri.lot_id
           WHERE r.reception_date=? ORDER BY c.name, p.description""",
        (date,),
    )
    return respond(rows, f"recepcion_{date}.xlsx")


@bp.get("/reports/dispatches-daily")
@login_required
def dispatches_daily():
    date = request.args.get("date", datetime.date.today().isoformat())
    rows = q(
        """SELECT d.dispatch_number, c.name as cliente, d.dispatch_date as fecha, d.guide_number as guia,
                  d.destination as destino, p.sku_code as sku, p.description as producto, l.lot_code as lote,
                  di.qty_requested as cantidad_solicitada, di.qty_conforming as conforme, di.qty_defective as defectuoso
           FROM dispatch_items di
           JOIN dispatches d ON d.id = di.dispatch_id
           JOIN clients c ON c.id = d.client_id
           JOIN products p ON p.id = di.product_id
           LEFT JOIN lots l ON l.id = di.lot_id
           WHERE d.dispatch_date=? ORDER BY c.name, p.description""",
        (date,),
    )
    return respond(rows, f"despacho_{date}.xlsx")


@bp.get("/reports/stock-by-client")
@login_required
def stock_by_client():
    rows = q(
        """SELECT c.name as cliente, p.sku_code as sku, p.description as producto, b.status as estado, SUM(b.qty) as cantidad
           FROM inventory_balances b JOIN products p ON p.id=b.product_id JOIN clients c ON c.id=b.client_id
           GROUP BY c.id, p.id, b.status ORDER BY c.name, p.description"""
    )
    return respond(rows, "stock_por_cliente.xlsx")


@bp.get("/reports/stock-by-product")
@login_required
def stock_by_product():
    rows = q(
        """SELECT p.sku_code as sku, p.description as producto, c.name as cliente, b.status as estado, SUM(b.qty) as cantidad
           FROM inventory_balances b JOIN products p ON p.id=b.product_id JOIN clients c ON c.id=b.client_id
           GROUP BY p.id, b.status ORDER BY p.description"""
    )
    return respond(rows, "stock_por_producto.xlsx")


@bp.get("/reports/stock-by-lot")
@login_required
def stock_by_lot():
    rows = q(
        """SELECT l.lot_code as lote, p.sku_code as sku, p.description as producto, c.name as cliente,
                  l.expiration_date as vencimiento, b.status as estado, SUM(b.qty) as cantidad
           FROM inventory_balances b
           JOIN products p ON p.id=b.product_id JOIN clients c ON c.id=b.client_id
           LEFT JOIN lots l ON l.id=b.lot_id
           WHERE b.lot_id IS NOT NULL
           GROUP BY l.id, b.status ORDER BY l.expiration_date"""
    )
    return respond(rows, "stock_por_lote.xlsx")


@bp.get("/reports/stock-by-location")
@login_required
def stock_by_location():
    rows = q(
        """SELECT loc.full_code as ubicacion, p.sku_code as sku, p.description as producto, c.name as cliente,
                  l.lot_code as lote, b.status as estado, b.qty as cantidad
           FROM inventory_balances b
           JOIN products p ON p.id=b.product_id JOIN clients c ON c.id=b.client_id
           LEFT JOIN lots l ON l.id=b.lot_id
           LEFT JOIN locations loc ON loc.id=b.location_id
           ORDER BY loc.full_code"""
    )
    return respond(rows, "stock_por_ubicacion.xlsx")


@bp.get("/reports/expirations")
@login_required
def expirations_report():
    days = int(request.args.get("days", 90))
    limit = (datetime.date.today() + datetime.timedelta(days=days)).isoformat()
    rows = q(
        """SELECT p.sku_code as sku, p.description as producto, c.name as cliente, l.lot_code as lote,
                  l.expiration_date as vencimiento, loc.full_code as ubicacion, b.qty as cantidad,
                  CAST(julianday(l.expiration_date) - julianday('now') AS INTEGER) as dias_restantes
           FROM inventory_balances b
           JOIN products p ON p.id=b.product_id JOIN clients c ON c.id=b.client_id
           JOIN lots l ON l.id=b.lot_id
           LEFT JOIN locations loc ON loc.id=b.location_id
           WHERE b.status='DISPONIBLE' AND l.expiration_date IS NOT NULL AND l.expiration_date <= ?
           ORDER BY l.expiration_date""",
        (limit,),
    )
    return respond(rows, "vencimientos.xlsx")


@bp.get("/reports/defects")
@login_required
def defects_report():
    rows = q(
        """SELECT r.reception_number, c.name as cliente, p.sku_code as sku, p.description as producto,
                  dt.name as tipo_defecto, qid.qty as cantidad, qi.inspection_date as fecha
           FROM quality_inspection_defects qid
           JOIN quality_inspections qi ON qi.id = qid.quality_inspection_id
           JOIN defect_types dt ON dt.id = qid.defect_type_id
           LEFT JOIN reception_items ri ON ri.id = qi.reception_item_id
           LEFT JOIN receptions r ON r.id = ri.reception_id
           LEFT JOIN products p ON p.id = ri.product_id
           LEFT JOIN clients c ON c.id = r.client_id
           ORDER BY qi.inspection_date DESC"""
    )
    return respond(rows, "productos_defectuosos.xlsx")


@bp.get("/reports/movements")
@login_required
def movements_report():
    date_from = request.args.get("from", (datetime.date.today() - datetime.timedelta(days=30)).isoformat())
    date_to = request.args.get("to", datetime.date.today().isoformat())
    rows = q(
        """SELECT m.movement_date as fecha, m.movement_type as tipo, c.name as cliente, p.sku_code as sku,
                  p.description as producto, l.lot_code as lote, m.qty as cantidad,
                  fl.full_code as origen, tl.full_code as destino, u.name as usuario, m.reason as motivo
           FROM inventory_movements m
           JOIN products p ON p.id=m.product_id JOIN clients c ON c.id=m.client_id
           LEFT JOIN lots l ON l.id=m.lot_id
           LEFT JOIN locations fl ON fl.id=m.from_location_id
           LEFT JOIN locations tl ON tl.id=m.to_location_id
           LEFT JOIN users u ON u.id=m.user_id
           WHERE date(m.movement_date) BETWEEN ? AND ?
           ORDER BY m.movement_date DESC""",
        (date_from, date_to),
    )
    return respond(rows, "movimientos.xlsx")


@bp.get("/reports/occupancy")
@login_required
def occupancy_report():
    rows = q(
        """SELECT wh.code as almacen, z.code as zona, rk.code as rack, loc.full_code as ubicacion,
                  loc.capacity as capacidad, loc.status as estado,
                  COALESCE((SELECT SUM(b.qty) FROM inventory_balances b WHERE b.location_id=loc.id),0) as ocupado
           FROM locations loc
           JOIN rack_levels lvl ON lvl.id=loc.rack_level_id
           JOIN racks rk ON rk.id=lvl.rack_id
           JOIN zones z ON z.id=rk.zone_id
           JOIN warehouses wh ON wh.id=z.warehouse_id
           ORDER BY wh.code, z.code, rk.code, loc.full_code"""
    )
    return respond(rows, "ocupacion_almacen.xlsx")
