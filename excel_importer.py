"""
Importador de Excel historico (CRP / CDP) con validacion en dos pasos.
Nunca importa una fila con error critico de forma silenciosa: primero se
valida y se muestra un resumen, y solo tras confirmacion explicita se
escribe en la base de datos.
"""
import re
import io
import uuid
import datetime
import subprocess
import tempfile
import os

import openpyxl

from db import q1, execute
from services import stock_service

# Cache en memoria de lotes validados pendientes de confirmar (demo de un solo proceso)
_PENDING_BATCHES = {}

CRP_FIELD_RULES = [
    ("fecha_ingreso", ["ingreso"]),
    ("cliente", ["cliente"]),
    ("contenedor", ["contenedor"]),
    ("pedido", ["pedido"]),
    ("guia", ["guía", "guia"]),
    ("procedencia", ["procedencia"]),
    ("tipo_carga", ["carga"]),
    ("descripcion", ["descripci"]),
    ("codigo_promocion", ["promoci"]),
    ("codigo_producto", ["código\nde\nproducto", "codigo de producto", "código de producto"]),
    ("lote", ["lote"]),
    ("cantidad_cajas", ["cajas"]),
    ("cantidad_unidades", ["unidades"]),
    ("fecha_vencimiento", ["vencimiento"]),
    ("oc_sede", ["oc", "sede"]),
    ("actividad", ["actividad"]),
    ("observaciones", ["observaci"]),
]

CDP_FIELD_RULES = [
    ("fecha_despacho", ["despacho"]),
    ("hora", ["hora"]),
    ("cliente", ["cliente"]),
    ("guia", ["guía", "guia"]),
    ("destino", ["destino"]),
    ("descripcion", ["descripci"]),
    ("contenedor_procedencia", ["contenedor"]),
    ("codigo_promocion", ["promoci"]),
    ("codigo_producto", ["código\nde producto", "codigo de producto", "código de producto", "código\nde\nproducto"]),
    ("lote", ["lote\nde producto", "lote de producto", "lote\nde\nproducto"]),
    ("cantidad_cajas", ["cajas"]),
    ("cantidad_conforme", ["conforme"]),
    ("cantidad_defectuosa", ["defectuos"]),
    ("fecha_vencimiento", ["vencimiento"]),
    ("aceptado_o_copacker", ["aceptad", "copacker"]),
    ("actividad", ["actividad"]),
    ("observaciones", ["observaci"]),
]


def _norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def _classify_headers(header_row, rules):
    mapping = {}
    used = set()
    for idx, cell in enumerate(header_row):
        text = _norm(cell)
        if not text:
            continue
        for field, keywords in rules:
            if field in mapping:
                continue
            if any(kw in text for kw in keywords):
                mapping[field] = idx
                used.add(idx)
                break
    return mapping


def _find_header_row(ws, keyword, max_scan=15):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        texts = [_norm(c) for c in row if c]
        if any(keyword in t for t in texts) and len(texts) > 5:
            return i, list(row)
    return None, None


def _parse_date(value):
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime.datetime):
        return value.date().isoformat(), None
    if isinstance(value, datetime.date):
        return value.isoformat(), None
    s = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(s, fmt).date().isoformat(), None
        except ValueError:
            continue
    return None, f"Fecha no interpretable: '{s}'"


def _safe_str(value):
    """Convierte cualquier valor de celda (incluyendo datetime.time) a texto plano serializable en JSON."""
    if value is None or value == "":
        return None
    return str(value)


def _to_number(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        try:
            return float(str(value).replace(",", "."))
        except ValueError:
            return None


def maybe_convert_xlsb(path):
    """Si el archivo es .xlsb, intenta convertirlo a .xlsx con LibreOffice (si esta disponible)."""
    if not path.lower().endswith(".xlsb"):
        return path
    outdir = tempfile.mkdtemp(prefix="wms_import_")
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir, path],
            check=True, timeout=120, capture_output=True,
        )
    except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        raise ValueError(
            "No se pudo convertir el archivo .xlsb automaticamente (LibreOffice no disponible en este servidor). "
            "Por favor exporte el archivo como .xlsx desde Excel y vuelva a intentarlo."
        ) from e
    base = os.path.splitext(os.path.basename(path))[0]
    return os.path.join(outdir, base + ".xlsx")


def _resolve_client(name_raw, sheet_name):
    """Resuelve el nombre de cliente contra clients/client_aliases; usa el nombre de hoja como respaldo."""
    candidate = (str(name_raw).strip() if name_raw else "") or sheet_name.strip()
    if not candidate:
        return None, None, "Sin nombre de cliente ni hoja identificable"
    key = candidate.strip().upper()
    row = q1("SELECT client_id FROM client_aliases WHERE upper(alias_text)=?", (key,))
    if row:
        c = q1("SELECT * FROM clients WHERE id=?", (row["client_id"],))
        return c["id"], c["name"], None
    row = q1("SELECT * FROM clients WHERE upper(name)=? OR upper(code)=?", (key, key))
    if row:
        return row["id"], row["name"], None
    return None, candidate, None  # cliente nuevo por crear


def parse_workbook(path, batch_type):
    """Lee TODAS las hojas del workbook y devuelve filas crudas con su hoja de origen."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    keyword = "ingreso" if batch_type == "CRP" else "despacho"
    rules = CRP_FIELD_RULES if batch_type == "CRP" else CDP_FIELD_RULES
    raw_rows = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row_idx, header_row = _find_header_row(ws, keyword)
        if not header_row_idx:
            continue
        mapping = _classify_headers(header_row, rules)
        if "cliente" not in mapping and "descripcion" not in mapping:
            continue
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if row is None or all(c in (None, "") for c in row):
                continue
            get = lambda f: row[mapping[f]] if f in mapping and mapping[f] < len(row) else None
            raw_rows.append({"sheet": sheet_name, "mapping": mapping, "raw": {
                k: get(k) for k, _ in rules
            }})
    wb.close()
    return raw_rows


def validate_rows(raw_rows, batch_type):
    validated = []
    seen_keys = set()
    for i, item in enumerate(raw_rows):
        r = item["raw"]
        sheet = item["sheet"]
        messages = []
        status = "valid"

        client_id, client_name, cerr = _resolve_client(r.get("cliente"), sheet)
        if client_id is None:
            status = "warning"
            messages.append(f"Cliente nuevo a crear: '{client_name}'")

        desc = (r.get("descripcion") or "").strip() if r.get("descripcion") else ""
        sku = str(r.get("codigo_producto")).strip() if r.get("codigo_producto") not in (None, "") else ""
        if not desc and not sku:
            status = "error"
            messages.append("Sin descripcion ni codigo de producto: no se puede identificar el item")
        if not sku:
            status = "warning" if status != "error" else status
            messages.append("Sin codigo de producto: se generara un codigo automatico")

        date_field = "fecha_ingreso" if batch_type == "CRP" else "fecha_despacho"
        date_val, date_err = _parse_date(r.get(date_field))
        if date_err:
            status = "error"
            messages.append(date_err)

        exp_val, exp_err = _parse_date(r.get("fecha_vencimiento"))
        if exp_err:
            status = "warning" if status != "error" else status
            messages.append(f"Vencimiento ignorado: {exp_err}")

        cajas = _to_number(r.get("cantidad_cajas"))
        if batch_type == "CRP":
            unidades = _to_number(r.get("cantidad_unidades"))
            if (unidades is None or unidades <= 0) and (cajas is None or cajas <= 0):
                status = "error"
                messages.append("Sin cantidad de cajas ni de unidades")
        else:
            conforme = _to_number(r.get("cantidad_conforme"))
            defectuosa = _to_number(r.get("cantidad_defectuosa"))
            if conforme is None and defectuosa is None:
                if cajas is None or cajas <= 0:
                    status = "error"
                    messages.append("Sin cantidad conforme, defectuosa ni cajas")
                else:
                    status = "warning" if status != "error" else status
                    messages.append("Sin cantidad en unidades: se usara la cantidad de cajas como estimado")

        if not r.get("lote"):
            status = "warning" if status != "error" else status
            messages.append("Sin lote registrado")

        dup_key = (sheet, client_name, date_val, r.get("guia"), sku or desc, r.get("lote"))
        is_dup = dup_key in seen_keys
        if is_dup:
            status = "duplicate"
            messages.append("Fila duplicada dentro del mismo archivo")
        seen_keys.add(dup_key)

        validated.append({
            "row_index": i, "sheet": sheet, "status": status, "messages": messages,
            "client_id": client_id, "client_name": client_name,
            "sku": sku, "description": desc, "lot_code": str(r.get("lote")).strip() if r.get("lote") else None,
            "date": date_val, "expiration_date": exp_val,
            "qty_cases": cajas,
            "qty_units": _to_number(r.get("cantidad_unidades")) if batch_type == "CRP" else None,
            "qty_conforming": _to_number(r.get("cantidad_conforme")) if batch_type == "CDP" else None,
            "qty_defective": _to_number(r.get("cantidad_defectuosa")) if batch_type == "CDP" else 0,
            "guide_number": _safe_str(r.get("guia")), "container_number": _safe_str(r.get("contenedor") or r.get("contenedor_procedencia")),
            "order_number": _safe_str(r.get("pedido")), "origin": _safe_str(r.get("procedencia")), "destination": _safe_str(r.get("destino")),
            "cargo_type": _safe_str(r.get("tipo_carga")), "purchase_order": _safe_str(r.get("oc_sede")),
            "activity_text": _safe_str(r.get("actividad")), "promotion_code": _safe_str(r.get("codigo_promocion")),
            "notes": _safe_str(r.get("observaciones")), "hora": _safe_str(r.get("hora")),
            "extra_field_value": _safe_str(r.get("aceptado_o_copacker")),
        })
    return validated


def summarize(validated):
    counts = {"valid": 0, "warning": 0, "error": 0, "duplicate": 0}
    for v in validated:
        counts[v["status"]] += 1
    return counts


def stage_import(path, batch_type):
    real_path = maybe_convert_xlsb(path)
    raw_rows = parse_workbook(real_path, batch_type)
    validated = validate_rows(raw_rows, batch_type)
    token = uuid.uuid4().hex[:12]
    _PENDING_BATCHES[token] = {"batch_type": batch_type, "rows": validated, "source_file": os.path.basename(path)}
    return token, validated, summarize(validated)


def _get_or_create_client(name):
    row = q1("SELECT * FROM clients WHERE upper(name)=?", (name.upper(),))
    if row:
        return row["id"]
    code = re.sub(r"[^A-Z0-9]", "", name.upper())[:12] or f"CLI{uuid.uuid4().hex[:6]}"
    base_code = code
    n = 1
    while q1("SELECT id FROM clients WHERE code=?", (code,)):
        n += 1
        code = f"{base_code}{n}"
    cid = execute("INSERT INTO clients (code, name) VALUES (?,?)", (code, name))
    execute("INSERT INTO client_configs (client_id) VALUES (?)", (cid,))
    return cid


def _get_or_create_product(client_id, sku, description):
    if not sku:
        sku = f"AUTO-{abs(hash(description)) % 100000}"
    row = q1("SELECT * FROM products WHERE client_id=? AND sku_code=?", (client_id, sku))
    if row:
        return row["id"]
    return execute("INSERT INTO products (sku_code, description, client_id) VALUES (?,?,?)",
                    (sku, description or sku, client_id))


def _get_or_create_lot(product_id, lot_code, expiration_date):
    if not lot_code:
        return None
    row = q1("SELECT * FROM lots WHERE product_id=? AND lot_code=?", (product_id, lot_code))
    if row:
        return row["id"]
    return execute("INSERT INTO lots (product_id, lot_code, expiration_date) VALUES (?,?,?)",
                    (product_id, lot_code, expiration_date))


def confirm_import(token, user_id, skip_errors=True):
    batch = _PENDING_BATCHES.get(token)
    if not batch:
        raise ValueError("Token de importacion no encontrado o expirado")
    rows = [r for r in batch["rows"] if r["status"] != "error" and (r["status"] != "duplicate")]
    batch_type = batch["batch_type"]

    batch_id = execute(
        "INSERT INTO import_batches (source_file, batch_type, imported_by, valid_count, warning_count, error_count, duplicate_count) VALUES (?,?,?,?,?,?,?)",
        (batch["source_file"], batch_type, user_id,
         len([r for r in batch["rows"] if r["status"] == "valid"]),
         len([r for r in batch["rows"] if r["status"] == "warning"]),
         len([r for r in batch["rows"] if r["status"] == "error"]),
         len([r for r in batch["rows"] if r["status"] == "duplicate"])),
    )

    groups = {}
    for r in rows:
        if not r["client_id"]:
            r["client_id"] = _get_or_create_client(r["client_name"])
        key = (r["client_id"], r["date"], r.get("guide_number"), r.get("container_number") if batch_type == "CRP" else None)
        groups.setdefault(key, []).append(r)

    created_headers = 0
    created_items = 0
    for (client_id, date, guide, container), items in groups.items():
        if not date:
            continue
        if batch_type == "CRP":
            number = f"REC-IMP-{batch_id}-{created_headers+1:04d}"
            rid = execute(
                """INSERT INTO receptions (reception_number, client_id, reception_date, container_number,
                     order_number, guide_number, origin, cargo_type, purchase_order, status, created_by, source)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                (number, client_id, date, container, items[0].get("order_number"), guide, items[0].get("origin"),
                 items[0].get("cargo_type"), items[0].get("purchase_order"), "COMPLETADO", user_id, "IMPORT_EXCEL"),
            )
            created_headers += 1
            for it in items:
                pid = _get_or_create_product(client_id, it["sku"], it["description"])
                lot_id = _get_or_create_lot(pid, it["lot_code"], it["expiration_date"])
                qty_units = it["qty_units"] or it["qty_cases"] or 0
                item_id = execute(
                    """INSERT INTO reception_items (reception_id, product_id, lot_id, qty_cases, qty_units, notes, quality_status, storage_status)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (rid, pid, lot_id, it["qty_cases"], qty_units, it.get("notes"), "DISPONIBLE", "PENDIENTE"),
                )
                created_items += 1
                if qty_units > 0:
                    stock_service.record_movement(
                        "RECEPCION", pid, client_id, qty_units, user_id, lot_id=lot_id,
                        to_location_id=None, to_status="DISPONIBLE",
                        reference_type="IMPORT", reference_id=rid,
                        reason=f"Carga historica desde Excel (lote de importacion #{batch_id})",
                        allow_negative=True,
                    )
        else:
            number = f"DES-IMP-{batch_id}-{created_headers+1:04d}"
            did = execute(
                """INSERT INTO dispatches (dispatch_number, client_id, dispatch_date, guide_number, destination,
                     status, created_by, source, closed_at)
                   VALUES (?,?,?,?,?,?,?,?,datetime('now'))""",
                (number, client_id, date, guide, items[0].get("destination"), "CERRADO", user_id, "IMPORT_EXCEL"),
            )
            created_headers += 1
            for it in items:
                pid = _get_or_create_product(client_id, it["sku"], it["description"])
                lot_id = _get_or_create_lot(pid, it["lot_code"], it["expiration_date"])
                conforme = it["qty_conforming"] or 0
                defectuosa = it["qty_defective"] or 0
                qty_total = conforme + defectuosa
                if qty_total <= 0:
                    qty_total = it["qty_cases"] or 0
                    conforme = qty_total
                item_id = execute(
                    """INSERT INTO dispatch_items (dispatch_id, product_id, lot_id, origin_container, qty_cases,
                         qty_requested, qty_conforming, qty_defective, expiration_date, notes)
                       VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    (did, pid, lot_id, it.get("container_number"), it["qty_cases"], qty_total, conforme,
                     defectuosa, it["expiration_date"], it.get("notes")),
                )
                created_items += 1
                if qty_total > 0:
                    stock_service.record_movement(
                        "DESPACHO", pid, client_id, qty_total, user_id, lot_id=lot_id,
                        from_location_id=None, from_status="DISPONIBLE",
                        reference_type="IMPORT", reference_id=did,
                        reason=f"Carga historica desde Excel (lote de importacion #{batch_id})",
                        allow_negative=True,
                    )

    del _PENDING_BATCHES[token]
    return {"batch_id": batch_id, "headers_created": created_headers, "items_created": created_items}
